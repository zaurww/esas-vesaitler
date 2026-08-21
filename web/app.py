"""Local web UI on localhost. Stdlib http.server only, no framework.

The server holds no state: every request reloads the client folder and
recomputes from events (CLAUDE.md §2, §3).
"""

from __future__ import annotations

import hashlib
import json
import os
import sys
import threading
import urllib.request
import webbrowser
from datetime import date as dt_date, datetime
from decimal import Decimal
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from engine.calc import (  # noqa: E402
    MONTHS_AZ, CalcError, compute_year, rate_matrix,
)
from engine.excel import build_import_template, build_workbook  # noqa: E402
from engine.mutate import (  # noqa: E402
    ACTIONS, IMPORT_FIELDS, export_client, guess_columns, rows_of,
    suggest_inv_no, verify_all,
)
from engine import rates  # noqa: E402
from engine.rates import CATEGORIES, CATEGORY_BY_CODE, ENGINE_VERSION  # noqa: E402
from engine.storage import DataError, list_clients, load_client  # noqa: E402
from web.update import (  # noqa: E402
    apply_update, check_latest, clear_pending_verify, download,
    read_pending_verify, rollback_update,
)

INDEX = Path(__file__).resolve().parent / "index.html"
STATIC = Path(__file__).resolve().parent / "static"

# --- who is already running -------------------------------------------
#
# Two copies of this program on the SAME clients/ folder is a data-loss bug,
# not a nuisance. _STATE_LOCK below serialises requests inside one process
# (§8.0); across processes there is no lock at all, and
# engine.mutate.core.transaction backs up, writes, reloads and rolls back on
# its own view of the folder -- so one process's rollback silently erases the
# other's write, and its write lands on data read before that write. A lost
# write with no error is exactly the failure §2.1 exists to forbid.
#
# Two copies in two DIFFERENT folders are legitimate: separate client sets,
# nothing shared. So identity is the folder, never the port.

_APP_MARKER = "esas-vesaitler"


def _install_id(root: Path) -> str:
    """Stable short id for one installation folder.

    Hashed rather than sent as a path: the probe answers any local program
    that asks, and where a bookkeeper keeps client tax data is not something
    to hand out for free (§3). A hash still compares exactly, which is all a
    caller needs. normcase first -- the same folder can be spelled with
    different case on Windows ("D:\\..." vs "d:\\...") and would otherwise
    look like a second installation.
    """
    return hashlib.sha256(
        os.path.normcase(str(root)).encode("utf-8")).hexdigest()[:16]


def instance_marker() -> dict:
    """Answer to "who holds this port?" -- served at /api/instance."""
    return {"app": _APP_MARKER, "install": _install_id(ROOT),
            "version": ENGINE_VERSION, "pid": os.getpid()}


def _probe(port: int, timeout: float = 1.0) -> dict | None:
    """Ask whatever holds `port` whether it is one of us. None = it is not."""
    try:
        with urllib.request.urlopen(
                f"http://127.0.0.1:{port}/api/instance", timeout=timeout) as r:
            # Capped read: an unrelated local server may answer with anything
            # at all, and a probe must not become a way to feed this process a
            # gigabyte before it has even started.
            return json.loads(r.read(4096).decode("utf-8"))
    except Exception:
        # Nothing listening, not HTTP, not JSON, too slow -- all the same
        # answer to the caller. Broad like update.check_latest's and for the
        # same reason: this is a courtesy at startup, not a calculation.
        return None


def _same_installation(info: dict | None) -> bool:
    """Is `info` another copy of us serving the same folder as this one?"""
    return (isinstance(info, dict)
            and info.get("app") == _APP_MARKER
            and info.get("install") == _install_id(ROOT))


_STATE_LOCK = threading.RLock()
_closed_cache: set = set()
_has_opening = False

# Result of the post-update check (see _run_startup_verify below), set once
# when this process starts serving and read by /api/update-status. None
# means "nothing to report" -- either no update is pending, or this process
# has not finished its startup check yet.
_startup_verify: dict | None = None


def _run_startup_verify(root: Path) -> dict | None:
    """If an update is pending verification, run it now and report the result.

    Called once, right as this process becomes the one actually serving
    (serve(), after a successful bind) -- never from a handoff process that
    only opens a browser and exits, and never on every request, because
    verify_all recomputes every closed year of every client and that cost
    belongs at startup, not on the hot path.

    Loud either way, matching §2.1: a clean result is printed too, not just
    silence, because "did the update check anything at all" should not be a
    question the accountant has to take on faith.
    """
    pending = read_pending_verify(root)
    if pending is None:
        return None
    bad = verify_all(root)
    if not bad:
        print(f"  yeniləmə yoxlanıldı ({pending['from_version']} → "
              f"{pending['to_version']}): bağlı illər uyğundur")
        clear_pending_verify(root)
        return {"ok": True, "from_version": pending["from_version"],
                "to_version": pending["to_version"]}
    total = sum(len(v) for v in bad.values())
    print(f"  ⚠ yeniləmədən sonra {total} uyğunsuzluq tapıldı "
          f"({pending['from_version']} → {pending['to_version']}) -- bax UI-də")
    # The marker is deliberately left in place: the problem is not fixed by
    # having been noticed, and clearing it here would make a rollback offered
    # later (once, from the UI) look unprompted the next time this page loads.
    return {"ok": False, "mismatches": bad,
            "from_version": pending["from_version"],
            "to_version": pending["to_version"]}
# Card fields the calculation never reads -- counterparty, e-invoice, serial,
# the client's own group.
# One dict rather than one global per field: they are all "what the card says
# about itself", and three parallel lookups was two too many.
_card_meta: dict = {}


def m(x: Decimal) -> str:
    """Money as a plain string; the UI formats for display."""
    return f"{x:.2f}"


def rate(x: Decimal) -> str:
    """A rate at full precision. A rate is not money -- do not send it as such.

    Rounded to the qəpik like an amount, the small-entrepreneur ceiling
    25% x 1.5 = 0.375 reached the browser as "0.38". That is not a display
    nicety: 0.38 is ABOVE the ceiling it claims to be, and the rate form,
    prefilled from the same figure, offered it back as an election the engine
    then refused -- the program appearing to reject a number it had just
    printed itself.
    """
    return format(x.normalize(), "f")


def available_years(data) -> list[int]:
    """Years the client can be opened on.

    Every year the data mentions, and then forward to the current one.

    That second half is the point. Balances carry themselves forward (§6.1),
    so the year after the last recorded event is a perfectly good report --
    but it could not be reached: the picker offered only years that already
    had rows, and no row can be written into a year that cannot be opened.
    The ways out were to close a year that had not been filed, or to buy
    something dated next year. Locking is a separate and deliberate act
    (§6.2); it must not be the toll for turning the page.

    Forward only -- years before the earliest recorded one are not invented.
    A year with nothing set up yet is not an error either: the report says so
    and offers to set the status, which is what puts the first row in it.
    """
    years = {ob.year for ob in data.opening_balances}
    years |= {s.year for s in data.statuses}
    years |= {a.in_date.year for a in data.assets if a.in_date}
    years = {y for y in years if y >= data.start_year}
    if not years:
        return [data.start_year]
    years |= set(range(max(years), max(max(years), dt_date.today().year) + 1))
    return sorted(years)


def default_year(data, years: list[int]) -> int:
    """Open on a year that actually computes.

    The newest year present in the data is often a year nothing has been set
    up for yet -- buy an asset in January and that year exists before anyone
    has recorded the taxpayer status for it. Landing there greets the user
    with a red error on a working installation, so prefer the newest year
    that has a status row.
    """
    ready = [s.year for s in data.statuses if s.year in years]
    return max(ready) if ready else (years[-1] if years else data.start_year)


def serialize(r) -> dict:
    return {
        "client_name": r.client_name,
        "voen": r.voen,
        "slug": r.slug,
        "start_year": r.start_year,
        "year": r.year,
        "is_closed": r.is_closed,
        "status": r.status,
        "status_name": r.status_name,
        "use_coefficient": r.use_coefficient,
        # Every status's coefficient for this year, so the status form can
        # show what each choice would raise the ceiling to. Sent from the
        # table rather than restated in the page: these are figures of the
        # law, and §5.1-bis keeps those out of code -- including this code.
        "coefficients": {s: {"name": n,
                             "coefficient": rate(rates.multiplier(r.year, s).coefficient)}
                         for s, n in rates.STATUS_NAMES.items()},
        "carried_from_prev": r.carried_from_prev,
        "closed_years": sorted(_closed_cache),
        "prev_year_closed": (r.year - 1) in _closed_cache,
        "has_opening": _has_opening,
        "engine_version": r.engine_version,
        "format_version": r.format_version,
        "months": MONTHS_AZ,
        "totals": {k: m(v) for k, v in r.totals.items()},
        "disposal_gain": m(r.disposal_gain),
        "disposal_loss": m(r.disposal_loss),
        "disposals": [
            {"inv_no": c.inv_no, "name": c.name, "category": c.category,
             "date": c.disposal_date.isoformat() if c.disposal_date else "",
             "type": c.disposal_type, "proceeds": m(c.proceeds),
             "residual": m(c.disposed), "gain_loss": m(c.gain_loss)}
            for c in r.disposed_cards
        ],
        # The figures that leave this program for the profit return (§12.1).
        # Built in the engine, not assembled in the page: the console and the
        # workbook print the same list, and three copies would drift.
        "declaration": [
            {"article": ln.article, "label_az": ln.label_az,
             "amount": m(ln.amount), "effect": ln.effect}
            for ln in r.declaration
        ],
        "declaration_deducted": m(r.declaration_deducted),
        "declaration_income": m(r.declaration_income),
        "declaration_net": m(r.declaration_net),
        "monthly": [m(v) for v in r.monthly],
        "warnings": r.warnings,
        "open_questions": r.open_questions,
        "categories": [
            {
                "code": c.code,
                "name_az": c.name_az,
                "name_ru": c.name_ru,
                "mixed_rates": c.mixed_rates,
                "rate": {
                    "statutory_max": rate(c.rate.statutory_max),
                    "statutory_year": c.rate.statutory_year,
                    "coefficient": str(c.rate.coefficient),
                    "ceiling": rate(c.rate.ceiling),
                    "applied": rate(c.rate.applied),
                    "elected": c.rate.elected,
                    "below_ceiling": c.rate.below_ceiling,
                    "below_statutory": c.rate.below_statutory,
                    "coefficient_used": c.rate.coefficient_used,
                    "source": c.rate.source,
                    # Which schedule, and how long it runs. A straight-line
                    # figure is never printed as a bare percentage of the
                    # base: that base is not what it was multiplied by.
                    "method": c.rate.method,
                    "term_years": c.rate.term_years,
                    "per_card": c.rate.per_card,
                },
                "opening": m(c.opening),
                "acquisition": m(c.acquisition),
                "addition": m(c.addition),
                "disposed": m(c.disposed),
                "depreciation": m(c.depreciation),
                "writeoff": m(c.writeoff),
                "closing": m(c.closing),
                "repair_limit": m(c.repair_limit),
                "repair_limit_pct": rate(c.repair_limit_pct),
                "repair_actual": m(c.repair_actual),
                "repair_deductible": m(c.repair_deductible),
                "repair_capitalized": m(c.repair_capitalized),
                "monthly": [m(v) for v in c.monthly],
                "cards": [
                    {
                        "asset_id": k.asset_id,
                        "inv_no": k.inv_no,
                        "name": k.name,
                        "category": k.category,
                        "in_date": k.in_date.isoformat() if k.in_date else "",
                        "cost": m(k.cost),
                        "is_legacy_pool": k.is_legacy_pool,
                        "counterparty": _card_meta.get(k.asset_id, {})
                                                  .get("counterparty", ""),
                        "e_qaime": _card_meta.get(k.asset_id, {})
                                             .get("e_qaime", ""),
                        "serial_no": _card_meta.get(k.asset_id, {})
                                               .get("serial_no", ""),
                        # Reporting only -- no figure below depends on it
                        # (§13.1). The name rides along with the id because
                        # every reader of this is about to print it.
                        "group_id": _card_meta.get(k.asset_id, {})
                                              .get("group_id", ""),
                        "group": _card_meta.get(k.asset_id, {})
                                           .get("group", ""),
                        "opening": m(k.opening),
                        "opening_source": k.opening_source,
                        "acquisition": m(k.acquisition),
                        "addition": m(k.addition),
                        "cost_effective": m(k.cost_effective),
                        "repair_actual": m(k.repair_actual),
                        "repair_deductible": m(k.repair_deductible),
                        "repair_capitalized": m(k.repair_capitalized),
                        "disposed": m(k.disposed),
                        "base": m(k.base),
                        "rate": rate(k.rate),
                        "depreciation": m(k.depreciation),
                        "writeoff": m(k.writeoff),
                        "closing": m(k.closing),
                        "rate_info": {
                            "statutory_max": rate(k.rate_info.statutory_max),
                            "statutory_year": k.rate_info.statutory_year,
                            "coefficient": str(k.rate_info.coefficient),
                            "ceiling": rate(k.rate_info.ceiling),
                            "applied": rate(k.rate_info.applied),
                            "source": k.rate_info.source,
                            "below_ceiling": k.rate_info.below_ceiling,
                            "below_statutory": k.rate_info.below_statutory,
                            "coefficient_used": k.rate_info.coefficient_used,
                            "method": k.rate_info.method,
                            "term_years": k.rate_info.term_years,
                            "remaining_years": k.rate_info.remaining_years,
                        } if k.rate_info else None,
                        "retired": k.retired,
                        "retired_kind": k.retired_kind,
                        "retired_year": k.retired_year,
                        "threshold_hit": k.threshold_hit,
                        "threshold_reason": k.threshold_reason,
                        "threshold_next": k.threshold_next,
                        "threshold_next_reason": k.threshold_next_reason,
                        "written_off": k.written_off,
                        "disposal_type": k.disposal_type,
                        "disposal_date": k.disposal_date.isoformat() if k.disposal_date else "",
                        "proceeds": m(k.proceeds),
                        "gain_loss": m(k.gain_loss),
                        "monthly": [m(v) for v in k.monthly],
                    }
                    for k in c.cards
                ],
            }
            for c in r.categories
        ],
    }


def parse_upload(body: dict) -> dict:
    """Turn an uploaded .xlsx / .csv / .tsv into rows of plain strings.

    Excel is the realistic source: every client already keeps their assets in
    a workbook. Dates are read as dates and re-rendered as ISO so they do not
    depend on the sheet's display format.
    """
    import base64
    import io as _io

    name = str(body.get("name", "")).lower()
    raw = base64.b64decode(body.get("b64", ""))
    if name.endswith((".csv", ".tsv", ".txt")):
        text = raw.decode("utf-8-sig", "replace")
        first = text.splitlines()[0] if text.strip() else ""
        sep = "\t" if "\t" in first else ";" if ";" in first else ","
        rows = [ln.split(sep) for ln in text.splitlines() if ln.strip()]
    else:
        from openpyxl import load_workbook
        wb = load_workbook(_io.BytesIO(raw), data_only=True, read_only=True)
        ws = wb[body["sheet"]] if body.get("sheet") in wb.sheetnames else wb.worksheets[0]
        rows = []
        for r in ws.iter_rows(values_only=True):
            out = []
            for c in r:
                if c is None:
                    out.append("")
                elif isinstance(c, datetime):
                    out.append(c.date().isoformat())
                elif isinstance(c, dt_date):
                    out.append(c.isoformat())
                elif isinstance(c, float) and c == int(c):
                    out.append(str(int(c)))
                else:
                    out.append(str(c))
            rows.append(out)
        rows = [r for r in rows if any(str(x).strip() for x in r)]
        body["sheets"] = wb.sheetnames

    width = max((len(r) for r in rows), default=0)
    rows = [list(r) + [""] * (width - len(r)) for r in rows]
    header = rows[0] if rows else []
    return {"header": header, "rows": rows[1:],
            "guess": guess_columns(header),
            "fields": list(IMPORT_FIELDS),
            "sheets": body.get("sheets", [])}


def asset_history(root: Path, slug: str, asset_id: str) -> dict:
    """Every year of one asset's life, recomputed from the events.

    The "лицевой счёт" an accountant expects: acquisition, then each year's
    opening -> additions -> repairs -> depreciation -> closing, ending in
    disposal or the current residual. Nothing here is stored -- it is the same
    per-year pipeline run repeatedly, so the history can never drift from the
    reports (§2).
    """
    rates.refresh(ROOT)
    data = load_client(root, slug)
    asset = next((a for a in data.assets if a.asset_id == asset_id), None)
    if asset is None:
        raise DataError(f"ƏV tapılmadı: {asset_id}")

    years = {s.year for s in data.statuses}
    years |= {ob.year for ob in data.opening_balances if ob.asset_id == asset_id}
    if asset.in_date:
        years.add(asset.in_date.year)
    years = sorted(y for y in years if y >= data.start_year)

    closed = data.closed_years()
    rows, seen = [], False
    for y in years:
        try:
            res = compute_year(data, y)
        except CalcError:
            continue
        card = next((c for c in res.cards if c.asset_id == asset_id), None)
        if card is None:
            if seen:
                break                       # already left the books
            continue
        seen = True
        ri = card.rate_info
        rows.append({
            "year": y,
            "closed": y in closed,
            "opening": m(card.opening),
            "acquisition": m(card.acquisition),
            "addition": m(card.addition),
            "repair_actual": m(card.repair_actual),
            "repair_deductible": m(card.repair_deductible),
            "repair_capitalized": m(card.repair_capitalized),
            "disposed": m(card.disposed),
            "base": m(card.base),
            "rate": rate(card.rate),
            "rate_ceiling": rate(ri.ceiling) if ri else "0",
            "rate_statutory": rate(ri.statutory_max) if ri else "0",
            "coefficient": str(ri.coefficient) if ri else "1",
            "method": ri.method if ri else "azalan",
            "term_years": ri.term_years if ri else None,
            "remaining_years": ri.remaining_years if ri else None,
            "depreciation": m(card.depreciation),
            "writeoff": m(card.writeoff),
            "closing": m(card.closing),
            "cost_effective": m(card.cost_effective),
            "gross_end": m(card.gross_end),
            "accumulated_end": m(card.accumulated_end),
            "threshold_hit": card.threshold_hit,
            "threshold_reason": card.threshold_reason,
            "written_off": card.written_off,
            "disposal_type": card.disposal_type,
            "disposal_date": card.disposal_date.isoformat() if card.disposal_date else "",
            "proceeds": m(card.proceeds),
            "gain_loss": m(card.gain_loss),
            "monthly": [m(v) for v in card.monthly],
        })

    events = []
    if asset.in_date:
        events.append({"date": asset.in_date.isoformat(), "kind": "Daxilolma",
                       "amount": m(asset.cost), "note": asset.counterparty})
    for ob in data.opening_balances:
        if ob.asset_id == asset_id:
            events.append({"date": f"{ob.year}-01-01", "kind":
                           "Açılış qalığı" if ob.source == "onboarding"
                           else "İl bağlanışı", "amount": m(ob.residual),
                           "note": ob.source})
    for a in data.additions:
        if a.asset_id == asset_id:
            events.append({"date": a.date.isoformat() if a.date else str(a.year),
                           "kind": "Dəyər artımı", "amount": m(a.amount),
                           "note": a.note})
    for rp in data.repairs:
        if rp.asset_id == asset_id:
            events.append({"date": rp.date.isoformat() if rp.date else str(rp.year),
                           "kind": "Təmir", "amount": m(rp.amount), "note": rp.note})
    for d in data.disposals:
        if d.asset_id == asset_id:
            events.append({"date": d.date.isoformat() if d.date else "",
                           "kind": "Xaricetmə", "amount": m(d.proceeds),
                           "note": d.type})
    for w in data.writeoffs:
        if w.asset_id == asset_id:
            events.append({"date": f"{w.year}-12-31", "kind": "Silinmə 500/5%",
                           "amount": "", "note": w.reason})
    events.sort(key=lambda e: e["date"])

    return {
        "asset": {
            "asset_id": asset.asset_id, "inv_no": asset.inv_no, "name": asset.name,
            "category": asset.category,
            "category_name": CATEGORY_BY_CODE[asset.category].name_az,
            "in_date": asset.in_date.isoformat() if asset.in_date else "",
            "cost": m(asset.cost), "counterparty": asset.counterparty,
            "e_qaime": asset.e_qaime, "serial_no": asset.serial_no,
            "is_legacy_pool": asset.is_legacy_pool, "note": asset.note,
        },
        "years": rows,
        "events": events,
        "months": MONTHS_AZ,
    }


def rate_report(root: Path, slug: str, first: int = 0, last: int = 0) -> dict:
    """Applied rates across years -- the cross-year control (§5.6.7).

    Every other report answers for one year. This one exists because a rate
    that is wrong is almost never wrong in a way a single year can show: it
    looks perfectly legal on its own and only stands out beside the years
    around it.
    """
    rates.refresh(ROOT)
    data = load_client(root, slug)
    years = available_years(data)
    if first:
        years = [y for y in years if y >= first]
    if last:
        years = [y for y in years if y <= last]
    mx = rate_matrix(data, years)

    def cell(c) -> dict:
        return {
            "year": c.year, "computed": c.computed, "on_books": c.on_books,
            "note": c.note,
            "statutory": rate(c.statutory), "coefficient": rate(c.coefficient),
            "ceiling": rate(c.ceiling), "applied": rate(c.applied),
            "source": c.source,
            "below_statutory": c.below_statutory,
            "below_ceiling": c.below_ceiling,
            "coefficient_used": c.coefficient_used,
            "law_changed": c.law_changed, "rate_changed": c.rate_changed,
            "cards": c.cards,
            "method": c.method, "term_years": c.term_years,
            "per_card": c.per_card,
        }

    def series(s) -> dict:
        return {
            "key": s.key, "kind": s.kind, "name": s.name,
            "subtitle": s.subtitle, "law_ref": s.law_ref,
            "category": s.category,
            "rate_changed": s.rate_changed, "law_changed": s.law_changed,
            "cells": [cell(c) for c in s.cells],
            "assets": [series(a) for a in s.assets],
        }

    return {
        "client": slug,
        "years": mx.years,
        "all_years": available_years(data),
        "closed": mx.closed,
        "failed": {str(k): v for k, v in mx.failed.items()},
        "rows": [series(s) for s in mx.rows],
    }


class Handler(BaseHTTPRequestHandler):
    server_version = "EsasVesaitler/" + ENGINE_VERSION

    def log_message(self, fmt, *args):  # keep the console quiet
        pass

    def _send(self, code: int, body: bytes, ctype: str, extra: dict | None = None):
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        for k, v in (extra or {}).items():
            self.send_header(k, v)
        self.end_headers()
        self.wfile.write(body)

    def _json(self, obj, code: int = 200):
        self._send(code, json.dumps(obj, ensure_ascii=False).encode("utf-8"),
                   "application/json; charset=utf-8")

    # One request at a time. The server is threaded, and two things it touches
    # are process-wide: the rate tables re-read by rates.refresh(), and the
    # per-request globals above that serialize() reads. Interleave two
    # requests and a report can be built from another request's client, or
    # computed while the rate table is being replaced.
    #
    # A lock rather than a redesign because the redesign is real work (pass
    # the rate table into the calculation instead of parking it in a module)
    # and this is a local single-user program: serialising requests costs
    # nothing here. A `compute_year` on a 5 000-card client takes under a
    # second, and nobody else is waiting.
    def do_GET(self):
        # The instance probe answers BEFORE the lock, deliberately. It reads
        # nothing mutable -- two constants and a pid -- and a copy that is
        # starting up needs the answer now: if this instance were mid-recompute
        # and holding the lock, a probe that blocked past its timeout would be
        # read as "not ours", and the newcomer would start a second writer on
        # this very folder. The check would fail exactly when it matters.
        if urlparse(self.path).path == "/api/instance":
            self._json(instance_marker())
            return
        with _STATE_LOCK:
            self._get()

    def do_POST(self):
        with _STATE_LOCK:
            self._post()

    # The page's own stylesheet and scripts. A whitelist of suffixes and a
    # containment check rather than "join the path and hope": this server
    # answers on localhost, but a request is still a string from outside, and
    # `/static/../../clients/x/config.toml` must not resolve to a file.
    STATIC_TYPES = {".js": "text/javascript; charset=utf-8",
                    ".css": "text/css; charset=utf-8"}

    def _static(self, name: str) -> None:
        path = (STATIC / name).resolve()
        if (path.suffix not in self.STATIC_TYPES
                or STATIC not in path.parents
                or not path.is_file()):
            self._json({"error": "not found"}, 404)
            return
        self._send(200, path.read_bytes(), self.STATIC_TYPES[path.suffix])

    def _get(self):
        url = urlparse(self.path)
        q = parse_qs(url.query)
        try:
            if url.path in ("/", "/index.html"):
                self._send(200, INDEX.read_bytes(), "text/html; charset=utf-8")
                return

            if url.path.startswith("/static/"):
                self._static(url.path[len("/static/"):])
                return

            if url.path == "/api/context":
                # Categories used to be fixed at import time, so which read
                # endpoint refreshed first never mattered. Now an owner can
                # append one (§12.4-quater), and this is the endpoint the
                # category dropdown is built from -- without a refresh here a
                # category added in an earlier session stayed invisible after
                # a restart until some unrelated action happened to trigger
                # one first.
                rates.refresh(ROOT)
                # One unreadable client must not blank the whole app: report it
                # per client so the picker still works and the message is seen.
                out = []
                for s in list_clients(ROOT):
                    try:
                        d = load_client(ROOT, s)
                    except DataError as e:
                        out.append({"slug": s, "name": s, "voen": "", "years": [],
                                    "closed_years": [], "error": str(e)})
                        continue
                    ys = available_years(d)
                    out.append({
                        "slug": s, "name": d.client_name, "voen": d.voen,
                        "years": ys,
                        "default_year": default_year(d, ys),
                        "closed_years": sorted(d.closed_years()),
                    })
                self._json({
                    "clients": out,
                    "engine_version": ENGINE_VERSION,
                    "categories": [
                        {"code": c.code, "name_az": c.name_az,
                         "name_ru": c.name_ru, "kind": c.kind}
                        for c in CATEGORIES
                    ],
                })
                return

            if url.path == "/api/report":
                slug = q.get("client", [""])[0]
                year = int(q.get("year", ["0"])[0])
                rates.refresh(ROOT)      # pick up edits without a restart
                data = load_client(ROOT, slug)
                global _closed_cache, _has_opening
                global _card_meta
                _card_meta = data.card_meta()
                _closed_cache = data.closed_years()
                _has_opening = any(ob.year == year for ob in data.opening_balances)
                if data.status_for(year) is None:
                    # Not a failure of the data -- a year nobody has set up
                    # yet. Say what is missing and let the UI offer the form.
                    self._json({
                        "error": f"{year} ili üçün sahibkarlıq statusu "
                                 f"göstərilməyib.",
                        "need": "status", "year": year, "kind": "setup",
                    }, 400)
                    return
                payload = serialize(compute_year(data, year))
                # The dictionary itself, not just the name on each card: the
                # filter and the card form need the whole list, including
                # groups nothing points at yet.
                payload["groups"] = [{"group_id": g.group_id, "name": g.name,
                                      "note": g.note} for g in data.groups]
                self._json(payload)
                return

            if url.path == "/api/rate-matrix":
                self._json(rate_report(
                    ROOT, q.get("client", [""])[0],
                    int(q.get("from", ["0"])[0] or 0),
                    int(q.get("to", ["0"])[0] or 0)))
                return

            if url.path == "/api/asset-history":
                self._json(asset_history(ROOT, q.get("client", [""])[0],
                                         q.get("asset_id", [""])[0]))
                return

            if url.path == "/api/import-template":
                blob = build_import_template(list(IMPORT_FIELDS))
                self._send(
                    200, blob,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    {"Content-Disposition":
                     'attachment; filename="EV-idxal-sablonu.xlsx"'},
                )
                return

            if url.path == "/api/client-archive":
                slug = q.get("client", [""])[0]
                blob = export_client(ROOT, slug)
                stamp = datetime.now().strftime("%Y%m%d")
                self._send(
                    200, blob, "application/zip",
                    {"Content-Disposition":
                     f'attachment; filename="{slug}-{stamp}.evbaza.zip"'},
                )
                return

            if url.path == "/api/next-inv":
                slug = q.get("client", [""])[0]
                cat = q.get("category", [""])[0]
                self._json({"inv_no": suggest_inv_no(
                    rows_of(ROOT, slug, "assets.tsv"), cat)})
                return

            if url.path == "/api/rates":
                year = int(q.get("year", ["0"])[0])
                rates.refresh(ROOT)
                self._json({"year": year,
                            "rates": rates.table_for(year),
                            "coefficients": rates.coefficients_for(year),
                            "parameters": rates.parameters_for(year),
                            "law_reviewed": rates.LAW_REVIEWED})
                return

            if url.path == "/api/export":
                slug = q.get("client", [""])[0]
                year = int(q.get("year", ["0"])[0])
                data = load_client(ROOT, slug)
                blob = build_workbook(compute_year(data, year),
                                      data.card_meta())
                fname = f"{slug}-{year}-amortizasiya.xlsx"
                self._send(
                    200, blob,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    {"Content-Disposition": f'attachment; filename="{fname}"'},
                )
                return

            if url.path == "/api/update-check":
                # Run once per page load (web/static/boot.js), never in a
                # background timer -- and check_latest() itself never raises,
                # so a worker offline just sees no banner, not an error here.
                self._json(check_latest())
                return

            if url.path == "/api/update-status":
                # The check already ran once, at this process's own startup
                # (serve() -> _run_startup_verify) -- this just hands the
                # result to whichever browser tab asks. {} rather than the
                # bare word "no": the frontend only has to test truthiness.
                self._json(_startup_verify or {})
                return

            self._json({"error": "not found"}, 404)

        except (DataError, CalcError) as e:
            # §2.1: a failed calculation is reported, never silently zeroed.
            self._json({"error": str(e), "kind": type(e).__name__}, 400)
        except Exception as e:  # noqa: BLE001
            self._json({"error": f"{type(e).__name__}: {e}", "kind": "internal"}, 500)

    def _post(self):
        """Every write goes through engine.mutate, which backs up, applies,
        re-validates the whole store and rolls back if it no longer parses."""
        url = urlparse(self.path)
        try:
            if url.path == "/api/parse-file":
                length = int(self.headers.get("Content-Length", "0"))
                body = json.loads(self.rfile.read(length) or b"{}")
                return self._json(parse_upload(body))
            if url.path == "/api/update-apply":
                # Not a client mutation -- no slug, no changelog, no per-year
                # recompute to verify against. Operates on ROOT itself, so it
                # does not go through ACTIONS (engine.mutate.core.transaction
                # is keyed by client, which this has none of).
                check = check_latest()
                if not check["available"]:
                    return self._json(
                        {"error": "Yenilənəcək versiya tapılmadı"}, 400)
                blob = download(check["zip_url"])
                message = apply_update(ROOT, blob, from_version=ENGINE_VERSION,
                                       to_version=check["latest"])
                return self._json({"ok": True, "message": message})
            if url.path == "/api/update-rollback":
                # backup_dir comes from the SERVER's own pending-verify
                # marker, never from the request body -- the same reasoning
                # as update-apply not trusting a client-sent zip_url: this
                # writes to the installation itself, so what gets restored
                # must be something this process already decided on, not
                # something a request asked for.
                pending = read_pending_verify(ROOT)
                if pending is None:
                    return self._json(
                        {"error": "Geri qaytarılacaq yeniləmə yoxdur"}, 400)
                backup_dir = ROOT / "backups" / "_app" / pending["backup_dir"]
                message = rollback_update(ROOT, backup_dir)
                global _startup_verify
                _startup_verify = None
                return self._json({"ok": True, "message": message})
            if url.path != "/api/action":
                return self._json({"error": "not found"}, 404)
            length = int(self.headers.get("Content-Length", "0"))
            body = json.loads(self.rfile.read(length) or b"{}")
            slug = body.get("client", "")
            action = body.get("action", "")
            if action not in ACTIONS:
                return self._json({"error": f"naməlum əməliyyat: {action}"}, 400)
            result = ACTIONS[action](ROOT, slug, body.get("payload") or {})
            self._json({"ok": True, "result": result})
        except (DataError, CalcError) as e:
            self._json({"error": str(e), "kind": type(e).__name__}, 400)
        except (KeyError, ValueError, TypeError) as e:
            self._json({"error": f"düzgün olmayan məlumat: {e}"}, 400)
        except Exception as e:  # noqa: BLE001
            self._json({"error": f"{type(e).__name__}: {e}", "kind": "internal"}, 500)


class LocalServer(ThreadingHTTPServer):
    """Loopback only -- this is a local tool, it must never be reachable from
    the network.

    allow_reuse_address is OFF on purpose. It reads like the fix for "port
    still in TIME_WAIT", but on Windows SO_REUSEADDR means something else
    entirely: a second process is allowed to bind a port another process is
    already listening on. Both instances then sit on 8777 and the OS hands
    each connection to whichever it likes -- so after an update the old copy
    keeps answering half the requests with the old engine. Measured here:
    a freshly started server never got a single request. Without it, bind
    fails cleanly and serve() steps up to the next port as intended.
    """
    allow_reuse_address = False
    daemon_threads = True


def _bind_or_handoff(port: int, tries: int, probe=_probe):
    """Bind a port -- or discover that this installation already holds one.

    Returns (server, port), or (None, port) meaning "we are already serving
    there; do not start a second one".

    Stepping up to the next free port survives, because a busy port is often
    somebody else entirely. What is new is that a busy port gets ASKED who it
    is first: an unrelated service or a second installation is stepped over as
    before, but our own folder is handed off to instead of duplicated.
    """
    for candidate in range(port, port + tries):
        try:
            return LocalServer(("127.0.0.1", candidate), Handler), candidate
        except OSError:
            if _same_installation(probe(candidate)):
                return None, candidate
            continue
    raise SystemExit(
        f"portlar {port}-{port + tries - 1} məşğuldur — "
        f"başqa bir proqram onları tutub?"
    )


def serve(port: int = 8777, open_browser: bool = True,
          tries: int = 12) -> None:
    """Start the local UI -- unless this installation is already serving.

    The step-up used to be unconditional, so a second double-click on
    Başlat.bat started a SECOND server on the same clients/ folder. That is
    a lost write waiting to happen, not a port nuisance -- see the note above
    _APP_MARKER. Now the newcomer hands off to the copy already running and
    stops, which is also what the user meant by double-clicking again: show
    me the program.
    """
    httpd, port = _bind_or_handoff(port, tries)
    url = f"http://127.0.0.1:{port}/"

    if httpd is None:
        print(f"Əsas Vəsaitlər artıq işləyir · {url}")
        print("  brauzerdə açılır — bu pəncərəni bağlaya bilərsiniz")
        print("  proqramı dayandırmaq üçün əvvəlki pəncərəni bağlayın")
        if open_browser:
            webbrowser.open(url)
        return

    global _startup_verify
    _startup_verify = _run_startup_verify(ROOT)

    print(f"Əsas Vəsaitlər · engine {ENGINE_VERSION}")
    print(f"  {url}")
    print("  dayandırmaq üçün: Ctrl+C")
    if open_browser:
        threading.Timer(0.6, lambda: webbrowser.open(url)).start()
    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        print("\ndayandırıldı")
    finally:
        httpd.server_close()


if __name__ == "__main__":
    serve()
